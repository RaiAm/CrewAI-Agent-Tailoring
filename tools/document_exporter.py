import docx
from docx.shared import Pt, Inches, RGBColor
import weasyprint
import markdown

class DocumentExporter:
    
    @staticmethod
    def export_to_docx(markdown_text: str, output_path: str = "tailored_resume.docx"):
        doc = docx.Document()
        
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)
            
        lines = markdown_text.split("\n")
        for line in lines:
            line = line.strip()
            if not line:
                continue
            if line.startswith("# "):
                p = doc.add_heading(line.replace("# ", ""), level=1)
                p.runs[0].font.color.rgb = RGBColor(0, 0, 0)
                p.runs[0].font.size = Pt(18)
            elif line.startswith("## "):
                p = doc.add_heading(line.replace("## ", ""), level=2)
                p.runs[0].font.color.rgb = RGBColor(30, 30, 30)
                p.runs[0].font.size = Pt(13)
            elif line.startswith("* ") or line.startswith("- "):
                doc.add_paragraph(line[2:], style='List Bullet')
            else:
                doc.add_paragraph(line)
                
        doc.save(output_path)
        return output_path

    @staticmethod
    def export_to_pdf(markdown_text: str, output_path: str = "tailored_resume.pdf"):
        html_content = markdown.markdown(markdown_text)
        
        styled_html = f"""
        <html>
        <head>
            <style>
                body {{
                    font-family: Arial, sans-serif;
                    margin: 40px;
                    line-height: 1.4;
                    color: #111;
                }}
                h1 {{ font-size: 22px; text-transform: uppercase; border-bottom: 2px solid #333; margin-bottom: 5px; }}
                h2 {{ font-size: 14px; text-transform: uppercase; border-bottom: 1px solid #ccc; margin-top: 15px; margin-bottom: 8px; }}
                p, li {{ font-size: 10.5pt; }}
                ul {{ margin-top: 3px; padding-left: 20px; }}
            </style>
        </head>
        <body>
            {html_content}
        </body>
        </html>
        """
        weasyprint.HTML(string=styled_html).write_pdf(output_path)
        return output_path